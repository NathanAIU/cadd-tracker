"""
export.py — Generate Excel and CSV downloads from the tools database.
"""

import io
import json
import datetime
import pandas as pd
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

DATA_DIR   = Path(__file__).parent.parent / "data"
TOOLS_FILE = DATA_DIR / "tools.json"

COLUMNS_DISPLAY = [
    "tool_name", "task", "category", "year", "authors",
    "doi", "journal", "github", "language", "open_source",
    "license", "molecule_type", "role", "key_features",
    "dependencies", "notes",
]

COLUMN_LABELS = {
    "tool_name":    "Tool Name",
    "task":         "Task / Application",
    "category":     "Category",
    "year":         "Year",
    "authors":      "Authors",
    "doi":          "DOI",
    "journal":      "Journal",
    "github":       "GitHub URL",
    "language":     "Language",
    "open_source":  "Open Source",
    "license":      "License",
    "molecule_type":"Molecule Type",
    "role":         "Role in Drug Discovery",
    "key_features": "Key Features",
    "dependencies": "Dependencies",
    "notes":        "Notes",
}

CATEGORY_COLORS = {
    "Molecular Docking":                  "D6E4F0",
    "Molecular Docking / Benchmarking":   "D6E4F0",
    "Molecular Docking / Validation":     "D6E4F0",
    "Molecular Docking / Representation Learning": "D6E4F0",
    "Generative Molecular Design":        "D5F5E3",
    "Generative Molecular Design / LLM":  "D5F5E3",
    "Property Prediction / QSAR":         "FDEBD0",
    "ADMET Prediction":                   "FADBD8",
    "3D-QSAR":                            "FADBD8",
    "Benchmark / Dataset":                "F0ECF8",
    "Dataset / Benchmark":                "F0ECF8",
    "Cheminformatics / Fingerprints":     "FEF9E7",
    "Cheminformatics":                    "FEF9E7",
    "Cheminformatics Platform":           "FEF9E7",
    "Free Energy Calculations":           "EAF7FB",
    "Molecular Dynamics Analysis":        "E8F8F5",
    "Molecular Dynamics":                 "E8F8F5",
    "Protein Structure Prediction":       "F5EEF8",
    "Protein Structure / Co-folding":     "F5EEF8",
    "Binding Site Prediction":            "FDEDEC",
    "PROTAC Design":                      "FEF5E7",
    "Linker Design":                      "FEF5E7",
    "Macrocycle / Cyclic Peptide":        "F9EBEA",
    "Macrocycle / Conformer Generation":  "F9EBEA",
    "Fragment-Based Drug Discovery":      "EBF5FB",
    "SAR / Lead Optimization":            "F4F6F7",
    "Synthesis Planning / Retrosynthesis":"E9F7EF",
    "Reaction Chemistry":                 "E9F7EF",
    "ML/DL Platform":                     "FDF2F8",
    "ML Platform / Benchmarking":         "FDF2F8",
    "Educational / Workflow":             "F7F9F9",
    "De Novo Design":                     "D5F5E3",
    "Pharmacophore / Virtual Screening":  "D6E4F0",
    "Drug-Target Interaction / KG":       "E8DAEF",
    "Molecular Representation / Pre-training": "FDFEFE",
    "Interaction Analysis":               "D6EAF8",
    "Molecular Design / Assembly":        "F0F3F4",
    "Molecular Preparation":              "D6E4F0",
}


def load_tools(status_filter="approved") -> list[dict]:
    if not TOOLS_FILE.exists():
        return []
    with open(TOOLS_FILE) as f:
        tools = json.load(f)
    if status_filter:
        tools = [t for t in tools if t.get("status", "approved") == status_filter]
    return tools


def tools_to_dataframe(tools: list[dict]) -> pd.DataFrame:
    rows = []
    for t in tools:
        row = {COLUMN_LABELS.get(c, c): t.get(c, "") for c in COLUMNS_DISPLAY}
        row["Open Source"] = "Yes" if t.get("open_source") else "No"
        rows.append(row)
    return pd.DataFrame(rows)


def to_csv(tools: list[dict]) -> bytes:
    df = tools_to_dataframe(tools)
    buf = io.StringIO()
    df.to_csv(buf, index=False)
    return buf.getvalue().encode("utf-8")


def to_excel(tools: list[dict]) -> bytes:
    if not HAS_OPENPYXL:
        # fallback to pandas xlsx
        df = tools_to_dataframe(tools)
        buf = io.BytesIO()
        df.to_excel(buf, index=False)
        return buf.getvalue()

    wb = Workbook()
    ws = wb.active
    ws.title = "CADD Tools"

    HEADER_FILL = PatternFill("solid", fgColor="1F3864")
    HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    DATA_FONT   = Font(name="Calibri", size=9)
    TOOL_FONT   = Font(name="Calibri", bold=True, size=9, color="1F3864")
    LINK_FONT   = Font(name="Calibri", size=9, color="2980B9", underline="single")
    thin  = Side(style="thin", color="CCCCCC")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    WRAP   = Alignment(wrap_text=True, vertical="top")
    CENTER = Alignment(horizontal="center", vertical="top")

    headers = [COLUMN_LABELS.get(c, c) for c in COLUMNS_DISPLAY]
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER

    for ri, tool in enumerate(tools, 2):
        cat  = tool.get("category", "")
        color = CATEGORY_COLORS.get(cat, "FAFAFA")
        fill  = PatternFill("solid", fgColor=color)

        for ci, col in enumerate(COLUMNS_DISPLAY, 1):
            val = tool.get(col, "")
            if col == "open_source":
                val = "Yes" if val else "No"
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.fill   = fill
            cell.border = BORDER
            cell.alignment = WRAP
            if col == "tool_name":
                cell.font = TOOL_FONT
            elif col in ("doi", "github"):
                cell.font = LINK_FONT
            elif col == "year":
                cell.alignment = CENTER
                cell.font = DATA_FONT
            else:
                cell.font = DATA_FONT

    col_widths = {
        "tool_name": 22, "task": 34, "category": 28, "year": 8,
        "authors": 22, "doi": 30, "journal": 32, "github": 40,
        "language": 14, "open_source": 10, "license": 14,
        "molecule_type": 22, "role": 30, "key_features": 50,
        "dependencies": 22, "notes": 45,
    }
    for ci, col in enumerate(COLUMNS_DISPLAY, 1):
        ws.column_dimensions[get_column_letter(ci)].width = col_widths.get(col, 18)

    ws.row_dimensions[1].height = 26
    for ri in range(2, len(tools) + 2):
        ws.row_dimensions[ri].height = 55

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS_DISPLAY))}1"

    # Stats sheet
    ws2 = wb.create_sheet("Stats")
    from collections import Counter
    cats = Counter(t.get("category", "Other") for t in tools)
    ws2.cell(1, 1, "Category").font = HEADER_FONT
    ws2.cell(1, 1).fill = HEADER_FILL
    ws2.cell(1, 2, "Count").font = HEADER_FONT
    ws2.cell(1, 2).fill = HEADER_FILL
    for ri, (cat, count) in enumerate(sorted(cats.items()), 2):
        ws2.cell(ri, 1, cat)
        ws2.cell(ri, 2, count)
    ws2.column_dimensions["A"].width = 40
    ws2.column_dimensions["B"].width = 10

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
